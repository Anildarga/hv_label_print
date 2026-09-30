from __future__ import annotations

import argparse
import ast
import json
import os
import re
import sys
from pathlib import Path
from typing import Any, Callable

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

DATABASE_URL_ENV = "DATABASE_URL"
MACHINE_HASH_PATTERN = re.compile(r"^[0-9a-fA-F]{64}$")
LEDGER_TABLE = "migration_imports"

DUMMY_SOURCE = _PROJECT_ROOT / "models" / "hv" / "data" / "scan_log.xlsx"
PACK_SOURCE = _PROJECT_ROOT / "models" / "hv" / "data" / "pack_id_scan_log.xlsx"
BMS_SOURCE = _PROJECT_ROOT / "models" / "hv" / "data" / "bms_scan_log.xlsx"
PACK_SESSION_SOURCE = _PROJECT_ROOT / "models" / "hv" / "data" / "pack_session.json"
USERS_SOURCE = _PROJECT_ROOT / "common" / "data" / "users.json"
MACHINE_LOCK_SOURCE = _PROJECT_ROOT / "common" / "auth" / "machine_lock.py"

SHEETS = {
    "dummy_log": (
        DUMMY_SOURCE,
        "HV Dummy",
        {
            "Sl.No": "sl_no",
            "Date": "date",
            "Time": "time",
            "Serial Count": "serial_count",
            "Pack ID (Serial Number)": "serial_number",
            "Ref No": "ref_no",
            "Variant": "variant",
            "M1": "m1",
            "M2": "m2",
            "M3": "m3",
            "M4": "m4",
            "Dummy Pack ID Status": "dummy_pack_status",
            "Dummy Pack QR Data": "dummy_pack_qr",
            "Issue Description 1": "issue_desc_1",
            "Issue Description 2": "issue_desc_2",
            "Action Plan": "action_plan",
            "Remark": "remark",
            "Admin Acknowledgement": "admin_ack",
        },
    ),
    "pack_id_log": (
        PACK_SOURCE,
        "Pack ID",
        {
            "Sl.No": "sl_no",
            "Date": "date",
            "Time": "time",
            "Serial Number": "serial_number",
            "Variant": "variant",
            "Pack QR Data": "pack_qr_data",
        },
    ),
    "bms_log": (
        BMS_SOURCE,
        "bmb_cmb",
        {
            "Sl.No": "sl_no",
            "Date": "date",
            "Time": "time",
            "Pack QR Data": "pack_qr_data",
            "BMB ID": "bmb_id",
            "CMB ID": "cmb_id",
            "Rework BMB ID": "rework_bmb_id",
            "Rework CMB ID": "rework_cmb_id",
        },
    ),
}

INSERT_COLUMNS = {
    "dummy_log": (
        "sl_no", "date", "time", "serial_count", "serial_number", "ref_no",
        "variant", "m1", "m2", "m3", "m4", "dummy_pack_status", "dummy_pack_qr",
        "issue_desc_1", "issue_desc_2", "action_plan", "remark", "admin_ack",
    ),
    "pack_id_log": ("sl_no", "date", "time", "serial_number", "variant", "pack_qr_data"),
    "bms_log": (
        "sl_no", "date", "time", "pack_qr_data", "bmb_id", "cmb_id",
        "rework_bmb_id", "rework_cmb_id",
    ),
}


class MigrationError(Exception):
    pass


def _read_authorized_hashes() -> list[str]:
    try:
        source = MACHINE_LOCK_SOURCE.read_text(encoding="utf-8")
        tree = ast.parse(source, filename=str(MACHINE_LOCK_SOURCE))
        raw_hashes: Any = None
        for node in tree.body:
            if isinstance(node, ast.Assign) and any(
                isinstance(target, ast.Name) and target.id == "AUTHORIZED_HASHES"
                for target in node.targets
            ):
                raw_hashes = ast.literal_eval(node.value)
                break
        if raw_hashes is None:
            raise MigrationError("AUTHORIZED_HASHES was not found in machine_lock.py.")
        if not isinstance(raw_hashes, (set, list, tuple)):
            raise MigrationError("AUTHORIZED_HASHES must be a set, list, or tuple.")
    except (OSError, SyntaxError, ValueError) as exc:
        raise MigrationError(f"Could not read authorized machine hashes: {exc}") from exc

    hashes: set[str] = set()
    for value in raw_hashes:
        if not isinstance(value, str) or not MACHINE_HASH_PATTERN.fullmatch(value):
            _log_skip("devices", f"Invalid machine hash in AUTHORIZED_HASHES: {value!r}")
            continue
        hashes.add(value.lower())
    return sorted(hashes)


def _log_skip(source: str, message: str) -> None:
    print(f"[SKIP] {source}: {message}", file=sys.stderr)


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def _excel_row(cursor: Any, table: str, device_id: Any, values: dict[str, Any]) -> None:
    columns = INSERT_COLUMNS[table]
    serial_number = values["sl_no"]
    if serial_number is None:
        columns = tuple(column for column in columns if column != "sl_no")
    elif isinstance(serial_number, bool) or not isinstance(serial_number, int):
        raise ValueError(f"Sl.No must be a whole number; got {serial_number!r}")
    supplied = ("device_id", *columns)
    placeholders = ", ".join(["%s"] * len(supplied))
    column_list = ", ".join(supplied)
    cursor.execute(
        f"INSERT INTO {table} ({column_list}) VALUES ({placeholders})",
        (device_id, *(values[column] for column in columns)),
    )


def _import_item(
    cursor: Any,
    source_key: str,
    source_label: str,
    source_name: str,
    insert: Callable[[], None],
    counts: dict[str, dict[str, int]],
) -> None:
    cursor.execute(
        f"SELECT status, reason FROM {LEDGER_TABLE} WHERE source_key = %s",
        (source_key,),
    )
    previous = cursor.fetchone()
    if previous is not None:
        counts[source_name]["skipped"] += 1
        return

    cursor.execute("SAVEPOINT migration_item")
    status = "imported"
    reason = ""
    try:
        insert()
    except Exception as exc:
        cursor.execute("ROLLBACK TO SAVEPOINT migration_item")
        status = "skipped"
        reason = f"{type(exc).__name__}: {exc}"
        counts[source_name]["skipped"] += 1
        _log_skip(source_label, reason)
    else:
        counts[source_name]["imported"] += 1
    finally:
        cursor.execute("RELEASE SAVEPOINT migration_item")

    cursor.execute(
        f"""
        INSERT INTO {LEDGER_TABLE} (source_key, status, reason)
        VALUES (%s, %s, %s)
        """,
        (source_key, status, reason),
    )


def _read_sheet(path: Path, sheet_name: str, header_map: dict[str, str]) -> list[tuple[int, dict[str, Any]]]:
    from openpyxl import load_workbook

    if not path.exists():
        raise MigrationError(f"Source file does not exist: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise MigrationError(f"Sheet {sheet_name!r} is missing from {path.name}.")
        worksheet = workbook[sheet_name]
        header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        header_indices = {
            str(value).strip(): index
            for index, value in enumerate(header_values)
            if value is not None
        }
        missing = sorted(set(header_map) - set(header_indices))
        if missing:
            raise MigrationError(
                f"{path.name}/{sheet_name} is missing columns: {', '.join(missing)}"
            )

        result = []
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            values: dict[str, Any] = {}
            for header, column in header_map.items():
                index = header_indices[header]
                value = row[index] if index < len(row) else None
                if column == "sl_no" and value not in (None, ""):
                    if isinstance(value, bool) or (
                        isinstance(value, float) and not value.is_integer()
                    ):
                        values[column] = value
                    else:
                        try:
                            values[column] = int(value)
                        except (ValueError, TypeError, OverflowError):
                            values[column] = value
                else:
                    values[column] = _cell_text(value)
            if any(value not in (None, "") for value in values.values()):
                result.append((row_number, values))
        return result
    finally:
        workbook.close()


def _import_sheet(
    cursor: Any,
    device_id: Any,
    table: str,
    path: Path,
    sheet_name: str,
    header_map: dict[str, str],
    counts: dict[str, dict[str, int]],
) -> None:
    source_name = path.name
    try:
        rows = _read_sheet(path, sheet_name, header_map)
    except Exception as exc:
        counts[source_name]["skipped"] += 1
        _log_skip(source_name, f"{type(exc).__name__}: {exc}")
        return

    for row_number, values in rows:
        source_key = f"excel:{path.name}:{sheet_name}:{row_number}"
        label = f"{path.name}/{sheet_name} row {row_number}"

        def insert_row(values: dict[str, Any] = values) -> None:
            _excel_row(cursor, table, device_id, values)

        _import_item(
            cursor, source_key, label, source_name, insert_row, counts
        )
    cursor.execute(
        f"""
        SELECT setval(
            pg_get_serial_sequence('{table}', 'sl_no'),
            GREATEST(COALESCE(MAX(sl_no), 1), 1),
            COALESCE(MAX(sl_no) >= 1, FALSE)
        )
        FROM {table}
        """
    )


def _import_devices(
    cursor: Any,
    hashes: list[str],
    counts: dict[str, dict[str, int]],
) -> dict[str, Any]:
    device_ids: dict[str, Any] = {}
    for machine_hash in hashes:
        source_key = f"authorized-device:{machine_hash}"
        existing = cursor.execute(
            "SELECT id FROM devices WHERE machine_hash = %s", (machine_hash,)
        ).fetchone()
        if existing is not None:
            device_ids[machine_hash] = existing[0]
            counts["AUTHORIZED_HASHES"]["skipped"] += 1
            continue

        def insert_device(machine_hash: str = machine_hash) -> None:
            cursor.execute(
                """
                INSERT INTO devices (machine_hash, name, is_active)
                VALUES (%s, %s, TRUE)
                RETURNING id
                """,
                (machine_hash, f"Imported device {machine_hash[:8]}"),
            )
            device_ids[machine_hash] = cursor.fetchone()[0]

        _import_item(
            cursor,
            source_key,
            f"machine_lock.py hash {machine_hash}",
            "AUTHORIZED_HASHES",
            insert_device,
            counts,
        )
        if machine_hash not in device_ids:
            existing = cursor.execute(
                "SELECT id FROM devices WHERE machine_hash = %s", (machine_hash,)
            ).fetchone()
            if existing is not None:
                device_ids[machine_hash] = existing[0]
    return device_ids


def _import_users(
    cursor: Any,
    path: Path,
    device_id: Any,
    counts: dict[str, dict[str, int]],
) -> None:
    source_name = path.name
    try:
        with path.open(encoding="utf-8") as file:
            users = json.load(file)
        if not isinstance(users, dict):
            raise MigrationError("users.json must contain a JSON object.")
    except Exception as exc:
        counts[source_name]["skipped"] += 1
        _log_skip(source_name, f"{type(exc).__name__}: {exc}")
        return

    for username, record in users.items():
        source_key = f"users.json:{username}"
        label = f"users.json user {username!r}"

        def insert_user(username: Any = username, record: Any = record) -> None:
            if not isinstance(username, str) or not username.strip():
                raise ValueError("username must be a non-empty string")
            if not isinstance(record, dict):
                raise ValueError("user entry must be an object")
            salt = record.get("salt")
            password_hash = record.get("password_hash")
            role = record.get("role", "operator")
            if not isinstance(salt, str) or not salt:
                raise ValueError("salt must be a non-empty string")
            if not isinstance(password_hash, str) or not password_hash:
                raise ValueError("password_hash must be a non-empty string")
            if role not in {"admin", "operator"}:
                raise ValueError(f"unsupported role {role!r}")
            cursor.execute(
                """
                INSERT INTO users
                    (device_id, username, password_salt, password_hash, role)
                VALUES (%s, %s, %s, %s, %s)
                """,
                (device_id, username, salt, password_hash, role),
            )

        _import_item(
            cursor, source_key, label, source_name, insert_user, counts
        )


def _import_pack_session(
    cursor: Any,
    path: Path,
    device_id: Any,
    counts: dict[str, dict[str, int]],
) -> None:
    source_name = path.name
    source_key = "pack_session.json:serial_counter"
    try:
        with path.open(encoding="utf-8") as file:
            state = json.load(file)
        if not isinstance(state, dict):
            raise MigrationError("pack_session.json must contain a JSON object.")
        serial_number = state.get("serial_number", "")
        serial_count = state.get("serial_count", "")
        if not isinstance(serial_number, str) or not serial_number.strip():
            raise MigrationError("serial_number is missing or invalid.")
        if not isinstance(serial_count, str) or not serial_count.strip():
            raise MigrationError("serial_count is missing or invalid.")
    except Exception as exc:
        counts[source_name]["skipped"] += 1
        _log_skip(source_name, f"{type(exc).__name__}: {exc}")
        return

    def insert_counter() -> None:
        cursor.execute(
            """
            INSERT INTO serial_counter (device_id, serial_number, serial_count)
            VALUES (%s, %s, %s)
            """,
            (device_id, serial_number.strip(), serial_count.strip()),
        )

    _import_item(
        cursor,
        source_key,
        source_name,
        source_name,
        insert_counter,
        counts,
    )


def migrate() -> dict[str, dict[str, int]]:
    database_url = os.getenv(DATABASE_URL_ENV, "").strip()
    if not database_url:
        raise MigrationError(f"{DATABASE_URL_ENV} is not set.")
    try:
        import psycopg
    except ImportError as exc:
        raise MigrationError(
            'Database driver is missing. Install it with: python -m pip install "psycopg[binary]"'
        ) from exc
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise MigrationError(
            "Excel reader is missing. Install it with: python -m pip install openpyxl"
        ) from exc

    hashes = _read_authorized_hashes()
    if not hashes:
        raise MigrationError("No valid machine hashes were found in AUTHORIZED_HASHES.")

    source_names = [path.name for path, _sheet, _mapping in SHEETS.values()]
    source_names.extend(
        [PACK_SESSION_SOURCE.name, USERS_SOURCE.name, "AUTHORIZED_HASHES"]
    )
    counts: dict[str, dict[str, int]] = {
        name: {"imported": 0, "skipped": 0} for name in source_names
    }
    try:
        with psycopg.connect(database_url) as connection, connection.cursor() as cursor:
            cursor.execute(
                f"""
                CREATE TABLE IF NOT EXISTS {LEDGER_TABLE} (
                    source_key TEXT PRIMARY KEY,
                    status TEXT NOT NULL CHECK (status IN ('imported', 'skipped')),
                    reason TEXT NOT NULL DEFAULT '',
                    created_at TIMESTAMPTZ NOT NULL DEFAULT now()
                )
                """
            )

            device_ids = _import_devices(cursor, hashes, counts)
            primary_device_id = device_ids.get(hashes[0])
            if primary_device_id is None:
                raise MigrationError("Could not resolve a device for imported records.")

            for table, (path, sheet, header_map) in SHEETS.items():
                _import_sheet(
                    cursor,
                    primary_device_id,
                    table,
                    path,
                    sheet,
                    header_map,
                    counts,
                )
            _import_pack_session(
                cursor, PACK_SESSION_SOURCE, primary_device_id, counts
            )
            _import_users(
                cursor, USERS_SOURCE, primary_device_id, counts
            )
    except MigrationError:
        raise
    except Exception as exc:
        raise MigrationError(f"Database migration failed: {exc}") from exc

    return dict(counts)


def main() -> int:
    argparse.ArgumentParser(
        description="Idempotently import local Excel/JSON data into PostgreSQL."
    ).parse_args()
    try:
        summary = migrate()
    except MigrationError as exc:
        print(f"Migration failed: {exc}", file=sys.stderr)
        return 1

    print("Migration summary (imported / skipped):")
    for source, totals in sorted(summary.items()):
        print(f"  {source}: {totals['imported']} imported, {totals['skipped']} skipped")
    print(
        f"  TOTAL: {sum(row['imported'] for row in summary.values())} imported, "
        f"{sum(row['skipped'] for row in summary.values())} skipped"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
