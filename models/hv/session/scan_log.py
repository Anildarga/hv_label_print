from __future__ import annotations

import os
import stat
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.protection import SheetProtection

from models.hv.session.paths import get_base_dir

LOG_FILE = get_base_dir() / "data" / "scan_log.xlsx"

SHEET_NAME = "HV Dummy"

PROTECT_SHEET = False #excel sheet protection is disables (False)


def _ensure_writable(path: Path) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.parent.chmod(stat.S_IRWXU | stat.S_IRWXG | stat.S_IRWXO)
        if path.exists():
            path.chmod(
                stat.S_IRUSR | stat.S_IWUSR |
                stat.S_IRGRP | stat.S_IWGRP |
                stat.S_IROTH | stat.S_IWOTH
            )
    except OSError:
        pass


COLUMNS = [
    "Sl.No",
    "Date",
    "Time",
    "Serial Count",
    "Pack ID (Serial Number)",
    "Ref No",
    "Variant",
    "M1",
    "M2",
    "M3",
    "M4",
    "Dummy Pack ID Status",
    "Dummy Pack QR Data",
    "Issue Description 1",
    "Issue Description 2",
    "Action Plan",
    "Remark",
    "Admin Acknowledgement",
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Segoe UI", bold=True, color="FFFFFF", size=10)
CELL_FONT   = Font(name="Segoe UI", size=10)
ALT_FILL    = PatternFill("solid", fgColor="EEF2F7")

COL_WIDTHS = {
    "Sl.No":                    7,
    "Date":                    13,
    "Time":                    12,
    "Serial Count":            12,
    "Ref No":                  12,
    "Variant":                 12,
    "M1":                      26,
    "M2":                      26,
    "M3":                      26,
    "M4":                      26,
    "Dummy Pack ID Status":    18,
    "Dummy Pack QR Data":      40,
    "Pack ID (Serial Number)": 22,
    "Issue Description 1":    30,
    "Issue Description 2":    30,
    "Action Plan":             30,
    "Remark":                  20,
    "Admin Acknowledgement":   30,
}

LEFT_ALIGN_COLS = {
    "M1", "M2", "M3", "M4",
    "Dummy Pack QR Data", "Pack ID (Serial Number)",
    "Issue Description 1", "Issue Description 2",
    "Action Plan", "Remark", "Admin Acknowledgement",
}

UNLOCKED_COLS = {
    "Ref No",
    "Serial Count",
    "Issue Description 1",
    "Issue Description 2",
    "Action Plan",
    "Remark",
    "Admin Acknowledgement",
}


@dataclass
class ScanRecord:
    date: str               = ""
    time: str               = ""
    serial_count: str       = ""
    ref_no: str             = ""
    variant: str            = ""
    m1: str                 = ""
    m2: str                 = ""
    m3: str                 = ""
    m4: str                 = ""
    dummy_pack_status: str  = ""
    dummy_pack_qr: str      = ""
    serial_number: str      = ""
    issue_desc_1: str       = ""
    issue_desc_2: str       = ""
    action_plan: str        = ""
    remark: str             = ""
    admin_ack: str          = ""

    @property
    def has_issue(self) -> bool:
        return bool(self.issue_desc_1.strip() or self.issue_desc_2.strip())

    @property
    def is_unacknowledged(self) -> bool:
        return self.has_issue and not self.admin_ack.strip()

    def as_row(self, sl_no: int) -> list:
        return [
            sl_no,
            self.date,
            self.time,
            self.serial_count,
            self.serial_number,
            self.ref_no,
            self.variant,
            self.m1,
            self.m2,
            self.m3,
            self.m4,
            self.dummy_pack_status,
            self.dummy_pack_qr,
            self.issue_desc_1,
            self.issue_desc_2,
            self.action_plan,
            self.remark,
            self.admin_ack,
        ]


def _protection(password: str, enabled: bool = True) -> SheetProtection:
    if not enabled:
        return SheetProtection(sheet=False)
    return SheetProtection(
        sheet=True,
        password=password,
        selectLockedCells=True,
        selectUnlockedCells=True,
        insertRows=False,
        insertColumns=False,
        deleteRows=False,
        deleteColumns=False,
        sort=False,
        autoFilter=False,
    )


def _build_header(ws: openpyxl.worksheet.worksheet.Worksheet, excel_password: str) -> None:
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS.get(col_name, 18)
        if col_name in UNLOCKED_COLS:
            cell.protection = openpyxl.styles.Protection(locked=False)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28
    ws.protection = _protection(excel_password, enabled=PROTECT_SHEET)


def _get_scan_log_sheet(wb: openpyxl.Workbook, excel_password: str) -> openpyxl.worksheet.worksheet.Worksheet:
    
    if SHEET_NAME in wb.sheetnames:
        return wb[SHEET_NAME]
    ws = wb.create_sheet(title=SHEET_NAME, index=0)
    _build_header(ws, excel_password)
    return ws


def _ensure_workbook(excel_password: str) -> openpyxl.Workbook:
    if LOG_FILE.exists():
        try:
            wb = openpyxl.load_workbook(LOG_FILE)
            ws = _get_scan_log_sheet(wb, excel_password)
            existing_headers = [
                str(ws.cell(row=1, column=i + 1).value or "").strip()
                for i in range(ws.max_column)
            ]
            if existing_headers != COLUMNS:
                data_rows = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if any(c is not None and str(c).strip() for c in row):
                        data_rows.append(list(row))
                old_col_map = {h: i for i, h in enumerate(existing_headers)}
                ws.protection = _protection(excel_password, enabled=False)
                ws.delete_rows(1, ws.max_row + 1)
                _build_header(ws, excel_password)
                ws.protection = _protection(excel_password, enabled=False)
                for row_idx, old_row in enumerate(data_rows, start=2):
                    alt = row_idx % 2 == 0
                    for col_idx, col_name in enumerate(COLUMNS, start=1):
                        old_i = old_col_map.get(col_name)
                        if old_i is not None and old_i < len(old_row):
                            value = old_row[old_i]
                        else:
                            value = ""
                        cell = ws.cell(row=row_idx, column=col_idx, value=value)
                        cell.font = CELL_FONT
                        h_align = "left" if col_name in LEFT_ALIGN_COLS else "center"
                        cell.alignment = Alignment(horizontal=h_align, vertical="center")
                        if alt:
                            cell.fill = ALT_FILL
                        if col_name in UNLOCKED_COLS:
                            cell.protection = openpyxl.styles.Protection(locked=False)
                    ws.row_dimensions[row_idx].height = 18
                ws.protection = _protection(excel_password, enabled=PROTECT_SHEET)
                _ensure_writable(LOG_FILE)
                wb.save(LOG_FILE)
            return wb
        except Exception:
            pass

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    _build_header(ws, excel_password)
    _ensure_writable(LOG_FILE)
    wb.save(LOG_FILE)
    return wb


def append_scan_record(record: ScanRecord, *, excel_password: str = "06082003") -> None:
    wb = _ensure_workbook(excel_password)
    ws = _get_scan_log_sheet(wb, excel_password)
    ws.protection = _protection(excel_password, enabled=False)

    sl_no = ws.max_row
    row_idx = ws.max_row + 1
    row_data = record.as_row(sl_no)

    for col_idx, (col_name, value) in enumerate(zip(COLUMNS, row_data), start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = CELL_FONT
        h_align = "left" if col_name in LEFT_ALIGN_COLS else "center"
        cell.alignment = Alignment(horizontal=h_align, vertical="center", wrap_text=True)
        if row_idx % 2 == 0:
            cell.fill = ALT_FILL
        if col_name in UNLOCKED_COLS:
            cell.protection = openpyxl.styles.Protection(locked=False)

    ws.row_dimensions[row_idx].height = 18
    ws.protection = _protection(excel_password, enabled=PROTECT_SHEET)
    _ensure_writable(LOG_FILE)
    wb.save(LOG_FILE)


def load_all_records_with_rows() -> list[tuple[int, ScanRecord]]:
    
    if not LOG_FILE.exists():
        return []
    try:
        wb = openpyxl.load_workbook(LOG_FILE, data_only=True)
        if SHEET_NAME not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[SHEET_NAME]
        records: list[tuple[int, ScanRecord]] = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            v = [str(c) if c is not None else "" for c in row]
            while len(v) < len(COLUMNS):
                v.append("")
            records.append((row_idx, ScanRecord(
                date=v[1],
                time=v[2],
                serial_count=v[3],
                serial_number=v[4],
                ref_no=v[5],
                variant=v[6],
                m1=v[7],
                m2=v[8],
                m3=v[9],
                m4=v[10],
                dummy_pack_status=v[11],
                dummy_pack_qr=v[12],
                issue_desc_1=v[13],
                issue_desc_2=v[14],
                action_plan=v[15],
                remark=v[16],
                admin_ack=v[17],
            )))
        wb.close()
        return records
    except Exception:
        return []


def load_all_records() -> list[ScanRecord]:
    return [rec for _row, rec in load_all_records_with_rows()]


def load_unacknowledged_issues() -> list[tuple[int, ScanRecord]]:
   
    return [(row, rec) for row, rec in load_all_records_with_rows() if rec.is_unacknowledged]


def patch_admin_ack(sheet_row: int, ack_text: str, *, excel_password: str = "06082003") -> bool:
    
    if not LOG_FILE.exists():
        return False
    wb = _ensure_workbook(excel_password)
    ws = _get_scan_log_sheet(wb, excel_password)
    if sheet_row < 2 or sheet_row > ws.max_row:
        wb.close()
        return False
    ws.protection = _protection(excel_password, enabled=False)
    col = COLUMNS.index("Admin Acknowledgement") + 1
    cell = ws.cell(row=sheet_row, column=col, value=ack_text)
    cell.font = CELL_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.protection = _protection(excel_password, enabled=PROTECT_SHEET)
    _ensure_writable(LOG_FILE)
    wb.save(LOG_FILE)
    return True


def add_sheet(sheet_name: str, *, excel_password: str = "06082003") -> tuple[bool, str]:
    if not sheet_name.strip():
        return False, "Sheet name cannot be empty."
    wb = _ensure_workbook(excel_password)
    existing = [ws.title.lower() for ws in wb.worksheets]
    if sheet_name.strip().lower() in existing:
        return False, f'Sheet "{sheet_name}" already exists.'
    new_ws = wb.create_sheet(title=sheet_name.strip())
    new_ws.sheet_properties.tabColor = "4472C4"
    _ensure_writable(LOG_FILE)
    wb.save(LOG_FILE)
    return True, f'Sheet "{sheet_name}" added.'


def open_log_file() -> None:
    if LOG_FILE.exists():
        os.startfile(str(LOG_FILE))