from __future__ import annotations

import os
import stat
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.protection import SheetProtection

from models.hv.session.paths import get_base_dir

LOG_FILE = get_base_dir() / "data" / "pack_id_scan_log.xlsx"
SHEET_NAME = "Pack ID"

PROTECT_SHEET = False

COLUMNS = ["Sl.No", "Date", "Time", "Serial Number", "Variant", "Pack QR Data"]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Segoe UI", bold=True, color="FFFFFF", size=10)
CELL_FONT   = Font(name="Segoe UI", size=10)
ALT_FILL    = PatternFill("solid", fgColor="EEF2F7")

COL_WIDTHS = {
    "Sl.No":          7,
    "Date":          13,
    "Time":          12,
    "Serial Number": 22,
    "Variant":       12,
    "Pack QR Data":  40,
}

LEFT_ALIGN_COLS = {"Serial Number", "Pack QR Data"}


@dataclass
class PackIdScanRecord:
    date: str          = ""
    time: str          = ""
    serial_number: str = ""
    variant: str        = ""
    pack_qr_data: str  = ""

    def as_row(self, sl_no: int) -> list:
        return [sl_no, self.date, self.time, self.serial_number, self.variant, self.pack_qr_data]


def _ensure_writable(path: Path) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.parent.chmod(stat.S_IRWXU | stat.S_IRWXG | stat.S_IRWXO)
        if path.exists():
            path.chmod(
                stat.S_IRUSR | stat.S_IWUSR |
                stat.S_IRGRP | stat.S_IWGRP |
                stat.S_IROTH | stat.S_IWOTH
            )
    except OSError:
        pass


def _protection(password: str, enabled: bool = True) -> SheetProtection:
    if not enabled:
        return SheetProtection(sheet=False)
    return SheetProtection(
        sheet=True, password=password,
        selectLockedCells=True, selectUnlockedCells=True,
        insertRows=False, insertColumns=False, deleteRows=False,
        deleteColumns=False, sort=False, autoFilter=False,
    )


def _build_header(ws, excel_password: str) -> None:
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS.get(col_name, 18)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28
    ws.protection = _protection(excel_password, enabled=PROTECT_SHEET)


def _get_sheet(wb: openpyxl.Workbook, excel_password: str):
    if SHEET_NAME in wb.sheetnames:
        return wb[SHEET_NAME]
    ws = wb.create_sheet(title=SHEET_NAME, index=0)
    _build_header(ws, excel_password)
    return ws


def _ensure_workbook(excel_password: str) -> openpyxl.Workbook:
    if LOG_FILE.exists():
        try:
            wb = openpyxl.load_workbook(LOG_FILE)
            _get_sheet(wb, excel_password)
            return wb
        except Exception:
            pass
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    _build_header(ws, excel_password)
    _ensure_writable(LOG_FILE)
    wb.save(LOG_FILE)
    return wb


def append_scan_record(record: PackIdScanRecord, *, excel_password: str = "06082003") -> None:
    wb = _ensure_workbook(excel_password)
    ws = _get_sheet(wb, excel_password)
    ws.protection = _protection(excel_password, enabled=False)

    sl_no = ws.max_row
    row_idx = ws.max_row + 1
    row_data = record.as_row(sl_no)
    for col_idx, (col_name, value) in enumerate(zip(COLUMNS, row_data), start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = CELL_FONT
        h_align = "left" if col_name in LEFT_ALIGN_COLS else "center"
        cell.alignment = Alignment(horizontal=h_align, vertical="center", wrap_text=True)
        if row_idx % 2 == 0:
            cell.fill = ALT_FILL

    ws.row_dimensions[row_idx].height = 18
    ws.protection = _protection(excel_password, enabled=PROTECT_SHEET)
    _ensure_writable(LOG_FILE)
    wb.save(LOG_FILE)


def load_all_records() -> list[PackIdScanRecord]:
    if not LOG_FILE.exists():
        return []
    try:
        wb = openpyxl.load_workbook(LOG_FILE, data_only=True)
        if SHEET_NAME not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[SHEET_NAME]
        records: list[PackIdScanRecord] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            v = [str(c) if c is not None else "" for c in row]
            while len(v) < len(COLUMNS):
                v.append("")
            records.append(PackIdScanRecord(
                date=v[1], time=v[2], serial_number=v[3], variant=v[4], pack_qr_data=v[5],
            ))
        wb.close()
        return records
    except Exception:
        return []


def open_log_file() -> None:
    if LOG_FILE.exists():
        os.startfile(str(LOG_FILE))